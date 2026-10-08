#!/usr/bin/env python3
"""Mini-bibliotheque de generation de documents : un meme contenu produit un .docx et un .pdf.

- DOCX : python-docx (styles Titre 1/2 reels, table des matieres a champ, pied de page numerote)
- PDF  : HTML + impression de Chrome/Edge headless (deux passes pour les numeros de page de la TOC)
"""
import html
import os
import re
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x1F, 0x4E, 0x79)
ACCENT_HEX = "1F4E79"
HEADER_FILL = "1F4E79"
ALT_FILL = "F2F5F9"

CHROME_CANDIDATES = [
    r"C:\Program Files\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    "/usr/bin/google-chrome", "/usr/bin/chromium", "/usr/bin/chromium-browser",
]


def find_browser():
    for c in CHROME_CANDIDATES:
        if os.path.exists(c):
            return c
    return shutil.which("chrome") or shutil.which("chromium")


INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def inline_html(text):
    out = []
    for part in INLINE.split(text):
        if part.startswith("**") and part.endswith("**"):
            out.append("<b>%s</b>" % html.escape(part[2:-2]))
        elif part.startswith("`") and part.endswith("`"):
            out.append("<code>%s</code>" % html.escape(part[1:-1]))
        else:
            out.append(html.escape(part))
    return "".join(out)


class Doc:
    def __init__(self, title, footer):
        self.title = title
        self.footer = footer
        self.blocks = []

    # ------------------------------------------------------------ construction
    def cover(self, **kw):
        self.blocks.append(("cover", kw))

    def toc(self, title="Table des matières"):
        self.blocks.append(("toc", title))

    def h1(self, text):
        self.blocks.append(("h1", text))

    def h2(self, text):
        self.blocks.append(("h2", text))

    def p(self, text, align="left", small=False):
        self.blocks.append(("p", text, align, small))

    def bullets(self, items):
        self.blocks.append(("bullets", items))

    def table(self, header, rows, widths=None, font=8.5, caption=None):
        self.blocks.append(("table", header, rows, widths, font, caption))

    def image(self, path, width_cm, caption=None):
        self.blocks.append(("image", path, width_cm, caption))

    def code(self, text, caption=None):
        self.blocks.append(("code", text, caption))

    def pagebreak(self):
        self.blocks.append(("pagebreak",))

    def landscape(self, on):
        self.blocks.append(("landscape", on))

    def note(self, text):
        self.blocks.append(("note", text))

    def headings(self):
        return [(b[0], b[1]) for b in self.blocks if b[0] in ("h1", "h2")]

    # ------------------------------------------------------------ DOCX
    def save_docx(self, path, page_numbers=None):
        page_numbers = page_numbers or {}
        d = Document()
        sec = d.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Cm(2.0))
        sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.9)

        st = d.styles
        st["Normal"].font.name = "Calibri"
        st["Normal"].font.size = Pt(10.5)
        st["Normal"].paragraph_format.space_after = Pt(5)
        st["Normal"].paragraph_format.line_spacing = 1.08
        for name, size, before, after in (("Heading 1", 15, 14, 6), ("Heading 2", 12, 9, 3)):
            s = st[name]
            s.font.name = "Calibri"
            s.font.size = Pt(size)
            s.font.bold = True
            s.font.color.rgb = ACCENT
            s.paragraph_format.space_before = Pt(before)
            s.paragraph_format.space_after = Pt(after)
            s.paragraph_format.keep_with_next = True
            rpr = s.element.get_or_add_rPr()
            rfonts = rpr.find(qn("w:rFonts"))
            if rfonts is None:
                rfonts = OxmlElement("w:rFonts")
                rpr.append(rfonts)
            for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(a), "Calibri")

        self._footer(sec)
        content_width = [Cm(17.0)]

        for b in self.blocks:
            kind = b[0]
            if kind == "cover":
                self._docx_cover(d, b[1])
            elif kind == "toc":
                self._docx_toc(d, b[1], page_numbers)
            elif kind == "h1":
                d.add_heading(b[1], level=1)
            elif kind == "h2":
                d.add_heading(b[1], level=2)
            elif kind == "p":
                para = d.add_paragraph()
                self._runs(para, b[1], size=9 if b[3] else None)
                para.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER,
                                  "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}[b[2]]
            elif kind == "bullets":
                for it in b[1]:
                    para = d.add_paragraph(style="List Bullet")
                    para.paragraph_format.space_after = Pt(2)
                    self._runs(para, it)
            elif kind == "table":
                self._docx_table(d, b[1], b[2], b[3], b[4], b[5], content_width[0])
            elif kind == "image":
                para = d.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.keep_with_next = bool(b[3])
                para.add_run().add_picture(b[1], width=Cm(b[2]))
                if b[3]:
                    self._caption(d, b[3])
            elif kind == "code":
                self._docx_code(d, b[1], b[2])
            elif kind == "note":
                self._docx_note(d, b[1])
            elif kind == "pagebreak":
                d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            elif kind == "landscape":
                new = d.add_section()
                if b[1]:
                    new.orientation = WD_ORIENT.LANDSCAPE
                    new.page_width, new.page_height = Cm(29.7), Cm(21)
                    new.left_margin = new.right_margin = Cm(1.5)
                    content_width[0] = Cm(26.7)
                else:
                    new.orientation = WD_ORIENT.PORTRAIT
                    new.page_width, new.page_height = Cm(21), Cm(29.7)
                    new.left_margin = new.right_margin = Cm(2.0)
                    content_width[0] = Cm(17.0)
                new.top_margin = new.bottom_margin = Cm(1.9)
        d.core_properties.title = self.title
        d.save(path)

    def _footer(self, sec):
        para = sec.footer.paragraphs[0]
        para.text = ""
        run = para.add_run(self.footer + "   |   Page ")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        self._field(para, "PAGE", size=8)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    @staticmethod
    def _field(para, instr, size=None, cached="1"):
        def r():
            run = para.add_run()
            if size:
                run.font.size = Pt(size)
                run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            return run
        a = r()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), "begin")
        a._r.append(e)
        b = r()
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = " %s " % instr
        b._r.append(t)
        c = r()
        e2 = OxmlElement("w:fldChar")
        e2.set(qn("w:fldCharType"), "separate")
        c._r.append(e2)
        x = r()
        x.text = cached
        y = r()
        e3 = OxmlElement("w:fldChar")
        e3.set(qn("w:fldCharType"), "end")
        y._r.append(e3)

    @staticmethod
    def _runs(para, text, size=None, color=None, bold=None):
        for part in INLINE.split(text):
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                run = para.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("`") and part.endswith("`"):
                run = para.add_run(part[1:-1])
                run.font.name = "Consolas"
                run.font.size = Pt((size or 10.5) - 1)
                continue
            else:
                run = para.add_run(part)
                if bold:
                    run.bold = True
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color

    @staticmethod
    def _shade(cell, fill):
        tcpr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcpr.append(shd)

    @staticmethod
    def _caption(d, text):
        para = d.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        run.italic = True
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x55, 0x5B, 0x66)

    def _docx_table(self, d, header, rows, widths, font, caption, total_width):
        if caption:
            para = d.add_paragraph()
            para.paragraph_format.keep_with_next = True
            run = para.add_run(caption)
            run.bold = True
            run.font.size = Pt(9)
        t = d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        total = sum(widths) if widths else len(header)
        ws = [total_width * (w / total) for w in (widths or [1] * len(header))]
        for i, h in enumerate(header):
            c = t.rows[0].cells[i]
            c.width = ws[i]
            self._shade(c, HEADER_FILL)
            c.paragraphs[0].paragraph_format.space_after = Pt(1)
            self._runs(c.paragraphs[0], h, size=font, color=RGBColor(255, 255, 255), bold=True)
        trpr = t.rows[0]._tr.get_or_add_trPr()
        hdr = OxmlElement("w:tblHeader")
        hdr.set(qn("w:val"), "true")
        trpr.append(hdr)
        for n, row in enumerate(rows):
            cells = t.add_row().cells
            cant = OxmlElement("w:cantSplit")
            t.rows[-1]._tr.get_or_add_trPr().append(cant)
            for i, val in enumerate(row):
                cells[i].width = ws[i]
                lines = str(val).split("\n")
                p0 = cells[i].paragraphs[0]
                p0.paragraph_format.space_after = Pt(1)
                p0.paragraph_format.line_spacing = 1.0
                self._runs(p0, lines[0], size=font)
                for ln in lines[1:]:
                    px = cells[i].add_paragraph()
                    px.paragraph_format.space_after = Pt(1)
                    px.paragraph_format.line_spacing = 1.0
                    self._runs(px, ln, size=font)
                if n % 2 == 1:
                    self._shade(cells[i], ALT_FILL)
        d.add_paragraph().paragraph_format.space_after = Pt(2)

    def _docx_code(self, d, text, caption):
        for ln in text.rstrip("\n").split("\n"):
            para = d.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.3)
            run = para.add_run(ln if ln else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(8)
            ppr = para._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F3F4F6")
            ppr.append(shd)
        if caption:
            self._caption(d, caption)
        else:
            d.add_paragraph().paragraph_format.space_after = Pt(2)

    def _docx_note(self, d, text):
        t = d.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.rows[0].cells[0]
        self._shade(c, "FFF7E6")
        c.paragraphs[0].paragraph_format.space_after = Pt(2)
        self._runs(c.paragraphs[0], text, size=9.5)
        d.add_paragraph().paragraph_format.space_after = Pt(2)

    def _docx_cover(self, d, kw):
        def line(text, size, bold=False, color=None, after=6, align=WD_ALIGN_PARAGRAPH.CENTER):
            para = d.add_paragraph()
            para.alignment = align
            para.paragraph_format.space_after = Pt(after)
            run = para.add_run(text)
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = color
        grey = RGBColor(0x55, 0x5B, 0x66)
        line(kw["university"], 15, True, ACCENT, 2)
        line(kw["faculty"], 12, False, grey, 2)
        line(kw["program"], 11, False, grey, 70)
        line(kw["doc_type"], 13, False, grey, 8)
        line(kw["title"], 26, True, ACCENT, 8)
        line(kw["subtitle"], 13, False, grey, 70)
        for label, value in kw["fields"]:
            para = d.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_after = Pt(3)
            r1 = para.add_run(label + " : ")
            r1.bold = True
            para.add_run(value)
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _docx_toc(self, d, title, page_numbers):
        para = d.add_paragraph()
        run = para.add_run(title)
        run.bold = True
        run.font.size = Pt(15)
        run.font.color.rgb = ACCENT
        entries = self.headings()
        first = True
        for i, (level, text) in enumerate(entries):
            para = d.add_paragraph()
            pf = para.paragraph_format
            pf.space_after = Pt(2 if level == "h2" else 3)
            pf.left_indent = Cm(0.6 if level == "h2" else 0)
            pf.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if first:
                for kind, txt in (("begin", None), ("instr", ' TOC \\o "1-2" \\h \\z \\u '), ("separate", None)):
                    r = para.add_run()
                    if kind == "instr":
                        it = OxmlElement("w:instrText")
                        it.set(qn("xml:space"), "preserve")
                        it.text = txt
                        r._r.append(it)
                    else:
                        fc = OxmlElement("w:fldChar")
                        fc.set(qn("w:fldCharType"), kind)
                        r._r.append(fc)
                first = False
            r = para.add_run("%s\t%s" % (text, page_numbers.get(text, "")))
            r.font.size = Pt(10.5 if level == "h1" else 10)
            r.bold = level == "h1"
            if i == len(entries) - 1:
                r2 = para.add_run()
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "end")
                r2._r.append(fc)
        d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------ HTML / PDF
    CSS = """
    @page { size: A4; margin: 1.9cm 2cm 2.1cm 2cm;
            @bottom-center { content: "%(footer)s   |   Page " counter(page); font: 8pt Calibri, 'Segoe UI', Arial, sans-serif; color: #6b7280; } }
    @page cover { margin: 1.9cm 2cm; @bottom-center { content: ""; } }
    @page land { size: A4 landscape; margin: 1.5cm 1.5cm 1.9cm 1.5cm;
            @bottom-center { content: "%(footer)s   |   Page " counter(page); font: 8pt Calibri, 'Segoe UI', Arial, sans-serif; color: #6b7280; } }
    * { box-sizing: border-box; }
    body { font-family: Calibri, 'Segoe UI', Arial, sans-serif; font-size: 10.5pt; line-height: 1.32; color: #1b1f24; margin: 0; }
    h1 { font-size: 15pt; color: #1f4e79; margin: 14pt 0 5pt; break-after: avoid; }
    h2 { font-size: 12pt; color: #1f4e79; margin: 9pt 0 3pt; break-after: avoid; }
    p { margin: 0 0 5pt; } p.center { text-align: center; } p.justify { text-align: justify; } p.small { font-size: 9pt; }
    ul { margin: 0 0 6pt; padding-left: 18pt; } li { margin-bottom: 2pt; }
    code { font-family: Consolas, 'Courier New', monospace; font-size: 0.9em; background: #f3f4f6; padding: 0 2px; }
    table { border-collapse: collapse; width: 100%%; margin: 2pt 0 8pt; }
    th { background: #1f4e79; color: #fff; text-align: left; font-weight: bold; }
    th, td { border: 0.6pt solid #9aa3ad; padding: 2.5pt 4pt; vertical-align: top; }
    tbody tr:nth-child(even) td { background: #f2f5f9; }
    tr { break-inside: avoid; } thead { display: table-header-group; }
    .tcap { font-weight: bold; font-size: 9pt; margin: 4pt 0 2pt; break-after: avoid; }
    figure { margin: 4pt 0 7pt; text-align: center; break-inside: avoid; }
    figure img { max-width: 100%%; } figcaption { font-style: italic; font-size: 8.5pt; color: #555b66; margin-top: 2pt; }
    pre { background: #f3f4f6; font-family: Consolas, 'Courier New', monospace; font-size: 8pt; line-height: 1.25;
          padding: 5pt 7pt; margin: 0 0 3pt; white-space: pre-wrap; break-inside: avoid; }
    .note { background: #fff7e6; border: 0.6pt solid #9aa3ad; padding: 4pt 6pt; margin: 3pt 0 8pt; font-size: 9.5pt; break-inside: avoid; }
    .cover { page: cover; text-align: center; break-after: page; min-height: 24cm; }
    .cover .u { font-size: 15pt; font-weight: bold; color: #1f4e79; margin-top: 0; }
    .cover .g { color: #555b66; }
    .cover .t { font-size: 26pt; font-weight: bold; color: #1f4e79; margin: 6pt 0; }
    .toc { break-after: page; } .toc h1 { margin-top: 0; }
    .toc div { display: flex; align-items: baseline; margin-bottom: 3pt; }
    .toc .l1 { font-weight: bold; } .toc .l2 { padding-left: 14pt; font-size: 10pt; }
    .toc .dots { flex: 1; border-bottom: 1px dotted #6b7280; margin: 0 4px; }
    .land { page: land; } .pb { break-after: page; }
    """

    def _html(self, page_numbers):
        out = ["<!doctype html><html lang='fr'><head><meta charset='utf-8'><title>%s</title><style>%s</style></head><body>"
               % (html.escape(self.title), self.CSS % {"footer": self.footer.replace('"', "'")})]
        land_open = False
        for b in self.blocks:
            k = b[0]
            if k == "cover":
                c = b[1]
                out.append("<div class='cover'><p class='u'>%s</p><p class='g'>%s</p><p class='g' style='margin-bottom:90pt'>%s</p>"
                           "<p class='g' style='font-size:13pt'>%s</p><p class='t'>%s</p><p class='g' style='font-size:13pt;margin-bottom:90pt'>%s</p>"
                           % tuple(html.escape(c[x]) for x in ("university", "faculty", "program", "doc_type", "title", "subtitle")))
                for label, value in c["fields"]:
                    out.append("<p><b>%s :</b> %s</p>" % (html.escape(label), html.escape(value)))
                out.append("</div>")
            elif k == "toc":
                out.append("<div class='toc'><h1>%s</h1>" % html.escape(b[1]))
                for level, text in self.headings():
                    out.append("<div class='%s'><span>%s</span><span class='dots'></span><span>%s</span></div>"
                               % ("l1" if level == "h1" else "l2", html.escape(text), page_numbers.get(text, "")))
                out.append("</div>")
            elif k in ("h1", "h2"):
                out.append("<%s>%s</%s>" % (k, html.escape(b[1]), k))
            elif k == "p":
                out.append("<p class='%s%s'>%s</p>" % (b[2], " small" if b[3] else "", inline_html(b[1])))
            elif k == "bullets":
                out.append("<ul>%s</ul>" % "".join("<li>%s</li>" % inline_html(i) for i in b[1]))
            elif k == "table":
                _, header, rows, widths, font, caption = b
                if caption:
                    out.append("<div class='tcap'>%s</div>" % html.escape(caption))
                total = sum(widths) if widths else len(header)
                cols = "".join("<col style='width:%.1f%%'>" % (100 * w / total) for w in (widths or [1] * len(header)))
                head = "".join("<th>%s</th>" % inline_html(h) for h in header)
                body = "".join("<tr>%s</tr>" % "".join("<td>%s</td>" % inline_html(str(v)).replace("\n", "<br>") for v in r) for r in rows)
                out.append("<table style='font-size:%.1fpt'><colgroup>%s</colgroup><thead><tr>%s</tr></thead><tbody>%s</tbody></table>"
                           % (font, cols, head, body))
            elif k == "image":
                src = "file:///" + os.path.abspath(b[1]).replace("\\", "/")
                cap = "<figcaption>%s</figcaption>" % html.escape(b[3]) if b[3] else ""
                out.append("<figure><img src='%s' style='width:%.1fcm'>%s</figure>" % (src, b[2], cap))
            elif k == "code":
                out.append("<pre>%s</pre>" % html.escape(b[1].rstrip("\n")))
                if b[2]:
                    out.append("<figure style='margin-top:0'><figcaption>%s</figcaption></figure>" % html.escape(b[2]))
            elif k == "note":
                out.append("<div class='note'>%s</div>" % inline_html(b[1]))
            elif k == "pagebreak":
                out.append("<div class='pb'></div>")
            elif k == "landscape":
                if b[1]:
                    out.append("<div class='land'>")
                    land_open = True
                elif land_open:
                    out.append("</div><div class='pb'></div>")
                    land_open = False
        if land_open:
            out.append("</div>")
        out.append("</body></html>")
        return "\n".join(out)

    def _render_pdf(self, html_text, pdf_path):
        browser = find_browser()
        if not browser:
            raise RuntimeError("Aucun Chrome/Edge trouve pour generer le PDF")
        with tempfile.TemporaryDirectory() as tmp:
            hp = os.path.join(tmp, "doc.html")
            with open(hp, "w", encoding="utf-8") as f:
                f.write(html_text)
            subprocess.run([browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                            "--allow-file-access-from-files", "--run-all-compositor-stages-before-draw",
                            "--virtual-time-budget=4000", "--print-to-pdf=" + os.path.abspath(pdf_path),
                            "file:///" + hp.replace("\\", "/")],
                           check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)

    def compute_pages(self, pdf_path):
        from pypdf import PdfReader
        reader = PdfReader(pdf_path)
        pages = [(p.extract_text() or "") for p in reader.pages]
        norm = lambda s: re.sub(r"\s+", " ", s).strip()
        result = {}
        toc_page = 1  # la TOC occupe la page 2 : on ne cherche les titres qu'apres
        for idx, text in enumerate(pages):
            if "Table des matières" in text and idx <= 2:
                toc_page = idx
        for level, heading in self.headings():
            for idx in range(toc_page + 1, len(pages)):
                if norm(heading) in norm(pages[idx]):
                    result[heading] = idx + 1
                    break
        return result, len(pages)

    def save_pdf(self, pdf_path):
        """Deux passes : la premiere mesure la pagination, la seconde ecrit les numeros dans la TOC."""
        self._render_pdf(self._html({}), pdf_path)
        numbers, count = self.compute_pages(pdf_path)
        self._render_pdf(self._html(numbers), pdf_path)
        numbers2, count2 = self.compute_pages(pdf_path)
        if numbers2 != numbers:
            self._render_pdf(self._html(numbers2), pdf_path)
            numbers, count = numbers2, count2
        return numbers, count
